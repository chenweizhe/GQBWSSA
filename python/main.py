import argparse
import json
import os
import sys
import time


if sys.platform == "win32":
    os.environ.setdefault("MKL_THREADING_LAYER", "TBB")

import numpy as np

from gqbwssa.algorithms import ALGORITHMS
from gqbwssa.data import load_dataset
from gqbwssa.fitness import FitnessRF

ROOT = os.path.dirname(os.path.abspath(__file__))

PLOT_STYLE = {
    "GWO": ("-^", "#FBEAA6"),
    "WOA": ("-d", "#D2D2D2"),
    "DBO": ("-s", "#9ACE87"),
    "HHO": ("-o", "#C76B60"),
    "GQBWSSA": ("-o", "#999BB9"),
}


def parse_args():
    p = argparse.ArgumentParser(description="GQBWSSA 特征选择（GPU 随机森林适应度）")
    p.add_argument("--data", default=os.path.join(ROOT, "data", "qifenlei_ct_quan.csv"))
    p.add_argument("--algos", default="GQBWSSA",
                   help="逗号分隔：GWO,WOA,DBO,HHO,GQBWSSA，或 all")
    p.add_argument("--runs", type=int, default=1, help="重复运行次数 Multi_Num")
    p.add_argument("--pop", type=int, default=20, help="种群规模 Popsize")
    p.add_argument("--iters", type=int, default=50, help="最大迭代次数 Max_iter")
    p.add_argument("--backend", choices=["gpu", "cpu"], default="gpu")
    p.add_argument("--ntree", type=int, default=20, help="随机森林决策树个数")
    p.add_argument("--mtry", type=int, default=2, help="每个节点的候选分裂特征数")
    p.add_argument("--bins", type=int, default=128, help="随机森林的特征分箱数")
    p.add_argument("--max-depth", type=int, default=16,
                   help="随机森林最大深度，0 表示不限制（更接近原来的完全生长，更慢）")
    p.add_argument("--streams", type=int, default=10, help="cuML 建树使用的 CUDA 流数")
    p.add_argument("--no-cache", action="store_true",
                   help="关闭适应度缓存，同一特征子集每次都重新训练")
    p.add_argument("--seed", type=int, default=None, help="随机种子（不设则每次不同，同 MATLAB）")
    p.add_argument("--out", default=os.path.join(ROOT, "results"))
    p.add_argument("--verbose", action="store_true", help="打印每次适应度评估")
    return p.parse_args()


def summarize(name, res1, res2, res3):
    final = res2[:, -1]
    ddof = 1 if len(final) > 1 else 0
    best_run = int(np.argmin(final))
    return {
        "name": name,
        "mean_acc": float(res1[0].mean()),
        "mean_ratio": float(res1[1].mean()),
        "mean_time": float(res1[2].mean()),
        "best_fit": float(final.min()),
        "worst_fit": float(final.max()),
        "std_fit": float(final.std(ddof=ddof)),
        "mean_fit": float(final.mean()),
        "best_pos": res3[best_run].astype(int).tolist(),
    }


def _setup_zh_font(plt):
    """Windows 用微软雅黑/黑体；Ubuntu 上这两款不在字体名里，改为直接加载字体文件。"""
    from matplotlib import font_manager

    candidates = [
        "/mnt/c/Windows/Fonts/msyh.ttc",
        "/mnt/c/Windows/Fonts/simhei.ttf",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            font_manager.fontManager.addfont(path)
            name = font_manager.FontProperties(fname=path).get_name()
            plt.rcParams["font.sans-serif"] = [name, "Microsoft YaHei", "SimHei", "DejaVu Sans"]
            break
    else:
        plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False


def plot_curves(curves, title, path):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    _setup_zh_font(plt)
    plt.figure(figsize=(8, 5))
    for name, curve in curves.items():
        marker, color = PLOT_STYLE.get(name, ("-o", None))
        plt.plot(np.arange(1, len(curve) + 1), curve, marker, color=color, linewidth=2,
                 markerfacecolor="#C76B60", markersize=3, label=f"{name}-RF")
    plt.grid(True)
    plt.legend(fontsize=12)
    plt.xlabel("进化代数", fontsize=12)
    plt.ylabel("适应度", fontsize=12)
    plt.title(title, fontsize=14)
    plt.tight_layout()
    plt.savefig(path, dpi=300)
    plt.close()


def append_excel(path, dataset_name, summaries):
    from openpyxl import Workbook, load_workbook

    wb = load_workbook(path) if os.path.exists(path) else Workbook()
    ws = wb.active
    order = "、".join(s["name"] for s in summaries)
    ws.append([dataset_name, "平均准确度", "平均特征选择率", "平均运行时间",
               "最佳适应度", "最差适应度", "标准差", "平均适应度", "最佳特征子集",
               f"按序号，优化算法依次为：{order}"])
    for k, s in enumerate(summaries, 1):
        ws.append([k, s["mean_acc"], s["mean_ratio"], s["mean_time"],
                   s["best_fit"], s["worst_fit"], s["std_fit"], s["mean_fit"], *s["best_pos"]])
    wb.save(path)


def main():
    args = parse_args()
    names = list(ALGORITHMS) if args.algos.lower() == "all" else [a.strip().upper() for a in args.algos.split(",")]
    for n in names:
        if n not in ALGORITHMS:
            raise SystemExit(f"Unknown algorithm {n}. Choices: {', '.join(ALGORITHMS)}")

    rng = np.random.default_rng(args.seed)
    dataset_name = os.path.basename(args.data)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    out_dir = os.path.join(args.out, f"{os.path.splitext(dataset_name)[0]}_{stamp}")
    os.makedirs(out_dir, exist_ok=True)

    t0 = time.perf_counter()
    xtrain, xvalid, ytrain, yvalid = load_dataset(args.data, rng)
    max_depth = None if args.max_depth <= 0 else args.max_depth
    fitness = FitnessRF(xtrain, ytrain, xvalid, yvalid, backend=args.backend, ntree=args.ntree,
                        mtry=args.mtry, n_bins=args.bins, max_depth=max_depth, n_streams=args.streams,
                        seed=args.seed, cache=not args.no_cache, verbose=args.verbose)
    print(f"Data ready: train {xtrain.shape}, test {xvalid.shape}, "
          f"backend={fitness.engine}, load {time.perf_counter() - t0:.1f}s", flush=True)

    dim = xtrain.shape[1]
    lb = -6 * np.ones(dim)
    ub = 6 * np.ones(dim)

    summaries, curves = [], {}
    for name in names:
        print(f"===== {name} start =====", flush=True)
        calls0, trains0, time0 = fitness.n_calls, fitness.n_trains, fitness.total_time
        res1, res2, res3 = ALGORITHMS[name](args.pop, args.iters, lb, ub, dim, fitness, args.runs,
                                            fitness, rng)
        s = summarize(name, res1, res2, res3)
        summaries.append(s)
        curves[name] = res2[0]
        n_calls = fitness.n_calls - calls0
        n_trains = fitness.n_trains - trains0
        print(f"{name}: ACC={s['mean_acc']:.5f}  feature_ratio={s['mean_ratio']:.4f}  "
              f"best_fitness={s['best_fit']:.6f}  worst_fitness={s['worst_fit']:.6f}  "
              f"mean_time={s['mean_time']:.1f}s  "
              f"RF calls {n_calls}, trained {n_trains}, "
              f"avg {(fitness.total_time - time0) / max(n_trains, 1):.3f}s/train", flush=True)

    with open(os.path.join(out_dir, "summary.json"), "w", encoding="utf-8") as f:
        json.dump({"args": vars(args), "results": summaries}, f, ensure_ascii=False, indent=2)
    plot_curves(curves, f"算法的迭代进化曲线({dataset_name})", os.path.join(out_dir, "convergence.png"))
    append_excel(os.path.join(args.out, "特征选择结果.xlsx"), dataset_name, summaries)
    print(f"Done. Results saved to {out_dir}", flush=True)


if __name__ == "__main__":
    main()
